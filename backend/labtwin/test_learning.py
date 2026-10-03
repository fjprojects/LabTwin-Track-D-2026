import io
import json
import tempfile
from pathlib import Path
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, override_settings
from django.urls import reverse

from .test_classrooms import ClassroomFixture
from .models import Course, CourseTopic, CourseMaterial, SourceChunk, LearningEvidence, TopicMastery, PracticeQuestion, Assignment, AssignmentAttempt, MonitoringEvent
from .learning.mastery import record_evidence


class LearningTests(ClassroomFixture):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.files = tempfile.TemporaryDirectory(prefix="labtwin-learning-test-")
        cls.learning_settings = override_settings(LABTWIN_MEDIA_ROOT=cls.files.name + "/uploads", LABTWIN_VECTOR_ROOT=cls.files.name + "/vectors", LABTWIN_VIDEO_OCR=False, PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
        cls.learning_settings.enable()

    @classmethod
    def tearDownClass(cls):
        cls.learning_settings.disable()
        super().tearDownClass()
        cls.files.cleanup()

    def setUp(self):
        super().setUp()
        self.join()
        self.course = Course.objects.create(classroom=self.room, name="Data Structures", language="C")
        self.topic = CourseTopic.objects.create(course=self.course, name="Linked Lists", description="Connect the new node to the previous head before updating the head pointer.")

    def upload(self, name="notes.txt", content=b"Linked lists: connect the new node to the previous head before updating the head pointer.", **kwargs):
        return self.tc.post(f"/api/learning/courses/{self.course.id}/materials/", {"file": SimpleUploadedFile(name, content), "topic_id": self.topic.id, **kwargs})

    def quiz(self):
        from .models import QuestionBankItem
        return QuestionBankItem.objects.create(topic=self.topic, title="Head update", prompt="Which update preserves the list?", options=["Connect old head", "Discard old head"], answer={"correct_index": 0})

    def question(self):
        self.quiz()
        return self.post(self.sc, f"learning/courses/{self.course.id}/practice/", {"topic_id": self.topic.id}).json()["question"]

    def test_auth_and_classroom_scope(self):
        self.assertEqual(Client().get("/api/learning/courses/").status_code, 401)
        self.assertEqual(self.oc.get(f"/api/learning/courses/{self.course.id}/materials/").status_code, 404)
        self.assertEqual(self.tc2.get(f"/api/learning/courses/{self.course.id}/materials/").status_code, 404)
        self.assertEqual(self.sc.get("/api/learning/courses/").json()["courses"][0]["id"], self.course.id)
        self.assertEqual(self.post(self.sc, "learning/courses/", {"classroom_id": self.room.id, "name": "Fake"}).status_code, 403)

    def test_text_vector_processing_and_citations(self):
        result = self.upload()
        self.assertEqual(result.status_code, 201, result.content)
        self.assertEqual(result.json()["material"]["status"], "ready", result.content)
        chunk = SourceChunk.objects.get()
        self.assertEqual(len(chunk.embedding), 384)
        self.assertEqual(chunk.metadata["classroom_id"], self.room.id)
        answer = self.post(self.sc, f"learning/courses/{self.course.id}/ask/", {"question": "Explain linked lists"}).json()
        self.assertTrue(answer["grounded"], answer)
        self.assertEqual(answer["citations"][0]["id"], chunk.id)
        self.assertEqual(answer["mode"], "source_excerpts")

    @override_settings(LABTWIN_PROCESS_INLINE=False)
    def test_material_worker_processes_the_queue_once(self):
        from django.core.management import call_command
        response = self.upload()
        material = CourseMaterial.objects.get(pk=response.json()["material"]["id"])
        self.assertEqual(material.status, "queued")
        self.assertFalse(SourceChunk.objects.exists())
        # The fixture's in-memory database cannot be shared with subprocesses.
        call_command("process_materials", in_process=True, stdout=io.StringIO())
        material.refresh_from_db()
        self.assertEqual(material.status, "ready")
        self.assertEqual(material.job.status, "ready")
        before = list(SourceChunk.objects.values_list("id", flat=True))
        call_command("process_materials", in_process=True, stdout=io.StringIO())
        self.assertEqual(list(SourceChunk.objects.values_list("id", flat=True)), before)
        self.assertEqual(material.job.attempts, 1)

    def test_lecture_summary_samples_the_whole_lecture(self):
        from .learning.ingestion import lecture_intelligence
        material = CourseMaterial(course=self.course, title="Long lecture", kind="video")
        units = [{"number": index + 1, "start_seconds": index * 60,
                  "text": f"Topic section {index}. This passage explains section {index}."}
                 for index in range(60)]
        with patch("labtwin.learning.ingestion.structured_completion", return_value=None) as completion:
            summary, chapters, _ = lecture_intelligence(material, units)
        self.assertIn("section 0", summary)
        self.assertIn("section 59", summary)
        self.assertEqual(len(completion.call_args.args[1]["passages"]), 12)
        self.assertGreater(chapters[-1]["start"], 3000)

    def test_pdf_page_and_pptx_slide_metadata(self):
        from reportlab.pdfgen import canvas
        from pptx import Presentation
        pdf = io.BytesIO()
        writer = canvas.Canvas(pdf)
        writer.drawString(50, 750, "Linked lists preserve the previous head.")
        writer.showPage(); writer.save()
        response = self.upload("notes.pdf", pdf.getvalue())
        self.assertEqual(response.json()["material"]["status"], "ready", response.content)
        self.assertEqual(SourceChunk.objects.first().unit.page_number, 1)
        slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = "Linked Lists"
        slide.placeholders[1].text = "Connect the new node to the old head."
        deck = io.BytesIO(); slides.save(deck)
        response = self.upload("lists.pptx", deck.getvalue())
        self.assertEqual(response.json()["material"]["status"], "ready", response.content)
        self.assertEqual(SourceChunk.objects.last().unit.slide_number, 1)
        self.assertTrue(SourceChunk.objects.last().unit.layout["shapes"])

    def test_corrupted_unsupported_and_encrypted_files_fail_safely(self):
        self.assertEqual(self.upload("payload.exe", b"not a material").status_code, 400)
        response = self.upload("broken.pdf", b"broken PDF")
        self.assertEqual(response.json()["material"]["status"], "failed")
        self.assertTrue(CourseMaterial.objects.get().file.storage.exists(CourseMaterial.objects.get().file.name))
        from pypdf import PdfWriter
        writer = PdfWriter(); writer.add_blank_page(200, 200); writer.encrypt("secret")
        pdf = io.BytesIO(); writer.write(pdf)
        response = self.upload("locked.pdf", pdf.getvalue())
        self.assertIn("Password", response.json()["material"]["error"])

    def test_vector_filter_prevents_cross_class_leakage(self):
        other_room = self.room.__class__.objects.create(teacher=self.teacher2, name="Other", join_code="OTHER98765432")
        other_course = Course.objects.create(classroom=other_room, name="Other course")
        response = self.tc2.post(f"/api/learning/courses/{other_course.id}/materials/", {"file": SimpleUploadedFile("secret.txt", b"Quantum entanglement secret from another classroom.")})
        self.assertEqual(response.json()["material"]["status"], "ready")
        self.upload()
        answer = self.post(self.sc, f"learning/courses/{self.course.id}/ask/", {"question": "Quantum entanglement"}).json()
        self.assertFalse(answer["grounded"])
        self.assertEqual(answer["citations"], [])

    def test_invalid_ai_citation_falls_back_to_real_excerpts(self):
        self.upload()
        with patch("labtwin.learning.rag.structured_completion", return_value={"supported": True, "claims": [{"text": "Invented fact", "quote": "this was invented", "source_id": 999999}]}):
            answer = self.post(self.sc, f"learning/courses/{self.course.id}/ask/", {"question": "Linked lists"}).json()
        self.assertEqual(answer["mode"], "source_excerpts")
        self.assertNotIn("Invented fact", answer["answer"])

    def test_source_media_range_and_revocation(self):
        self.upload()
        chunk = SourceChunk.objects.get()
        data = self.sc.get(f"/api/learning/sources/{chunk.id}/").json()
        url = data["media_url"].replace("http://testserver", "")
        response = Client().get(url, HTTP_RANGE="bytes=0-5")
        self.assertEqual(response.status_code, 206)
        self.assertEqual(b"".join(response.streaming_content), b"Linked")
        self.assertEqual(Client().get(url, HTTP_RANGE="bytes=999999-").status_code, 416)
        self.assertEqual(Client().get(f"/api/learning/materials/{chunk.unit.material_id}/media/").status_code, 401)
        self.tc.delete(f"/api/classrooms/{self.room.id}/students/{self.student.student_id}/")
        self.assertEqual(Client().get(url).status_code, 404)
        self.assertEqual(self.sc.get(f"/api/learning/sources/{chunk.id}/").status_code, 404)

    def test_expired_session_invalidates_signed_link(self):
        self.upload()
        chunk = SourceChunk.objects.get()
        url = self.sc.get(f"/api/learning/sources/{chunk.id}/").json()["media_url"].replace("http://testserver", "")
        self.post(self.sc, "auth/logout/", {})
        self.assertEqual(Client().get(url).status_code, 401)

    def test_lecture_timestamps_and_caption_validation(self):
        with patch("labtwin.learning.extraction.media_info", return_value=60):
            result = self.upload("lecture.mp4", b"placeholder-test-only", captions=json.dumps([{"start": 12, "end": 24, "text": "Linked lists connect the new node to the previous head."}]))
        self.assertEqual(result.json()["material"]["status"], "ready", result.content)
        chunk = SourceChunk.objects.get()
        self.assertEqual(chunk.unit.start_seconds, 12)
        data = self.sc.get(f"/api/learning/sources/{chunk.id}/").json()
        self.assertIn("0:12", data["source"]["label"])
        from .learning.transcription import validate_transcript
        with self.assertRaises(Exception):
            validate_transcript([{"start": 25, "end": 12, "text": "invalid"}], 60)

    def test_mastery_is_idempotent_and_recent_improvement_changes_path(self):
        student = self.student.student
        record_evidence(student, self.topic, "a", "programming", 0, misconception="Head pointer")
        record_evidence(student, self.topic, "a", "programming", 100)
        self.assertEqual(LearningEvidence.objects.count(), 1)
        record_evidence(student, self.topic, "b", "programming", 100)
        record_evidence(student, self.topic, "c", "quiz", 100)
        record_evidence(student, self.topic, "d", "quiz", 100)
        self.assertEqual(TopicMastery.objects.get().state, "Strong")
        path = self.sc.get(f"/api/learning/courses/{self.course.id}/path/").json()
        self.assertEqual(path["tasks"], [])

    def test_weak_prerequisites_reduce_difficulty(self):
        prerequisite = CourseTopic.objects.create(course=self.course, name="Pointers")
        self.topic.prerequisites.add(prerequisite)
        from .models import QuestionBankItem
        QuestionBankItem.objects.create(topic=prerequisite, title="Pointer", prompt="What is a pointer?", options=["Reference", "Loop"], answer={"correct_index": 0})
        response = self.post(self.sc, f"learning/courses/{self.course.id}/practice/", {"topic_id": self.topic.id})
        self.assertEqual(response.json()["question"]["topic_id"], prerequisite.id)
        self.assertIn("prerequisite", response.json()["question"]["reason"])

    def test_practice_answers_tests_and_solutions_are_private(self):
        question = self.question()
        for field in ("tests", "answer", "solution", "rubric"):
            self.assertNotIn(field, question)
        self.assertEqual(self.sc.get(f"/api/learning/questions/{question['id']}/solution/").status_code, 403)
        self.assertEqual(self.sc.get(f"/api/learning/topics/{self.topic.id}/bank/").status_code, 403)

    def test_quiz_server_grading_and_duplicate_request(self):
        question = self.question()
        import uuid
        data = {"request_id": str(uuid.uuid4()), "answer": "0", "score": 0}
        url = f"learning/questions/{question['id']}/attempts/"
        response = self.post(self.sc, url, data)
        self.assertEqual(response.json()["attempt"]["score"], 100)
        self.post(self.sc, url, data)
        self.assertEqual(LearningEvidence.objects.count(), 1)
        self.assertEqual(self.post(self.oc, url, data).status_code, 404)

    def test_code_results_hide_input_expected_and_stdout(self):
        from .models import QuestionBankItem
        QuestionBankItem.objects.create(topic=self.topic, title="Code", kind="code", prompt="Connect the nodes", tests=[{"input": "secret", "expected": "private"}])
        question = self.post(self.sc, f"learning/courses/{self.course.id}/practice/", {}).json()["question"]
        import uuid
        with patch.object(self.fake_views, "run_tests", return_value=(0, [{"passed": False, "input": "secret", "expected": "private", "stdout": "secret"}]), create=True):
            response = self.post(self.sc, f"learning/questions/{question['id']}/attempts/", {"request_id": str(uuid.uuid4()), "code": "print(input())"})
        self.assertNotIn("secret", response.content.decode())
        self.assertNotIn("private", response.content.decode())

    def test_runner_failure_saves_attempt_without_mastery_credit(self):
        from .models import QuestionBankItem, PracticeAttempt
        QuestionBankItem.objects.create(topic=self.topic, title="Code", kind="code", prompt="Code task", tests=[{"input": "", "expected": "1"}])
        question = self.post(self.sc, f"learning/courses/{self.course.id}/practice/", {}).json()["question"]
        import uuid
        with patch.object(self.fake_views, "run_tests", side_effect=RuntimeError("runner down"), create=True):
            response = self.post(self.sc, f"learning/questions/{question['id']}/attempts/", {"request_id": str(uuid.uuid4()), "code": "print(1)"})
        self.assertEqual(response.status_code, 503)
        self.assertEqual(PracticeAttempt.objects.get().code, "print(1)")
        self.assertEqual(LearningEvidence.objects.count(), 0)

    def test_hints_provide_progressive_help_without_full_solution(self):
        question = self.question()
        responses = [self.post(self.sc, "learning/hint/", {"question_id": question["id"], "level": level}).json() for level in (1, 2, 3)]
        self.assertIn("Pseudocode", responses[-1]["message"])
        self.assertEqual([r["level"] for r in responses], [1, 2, 3])

    def test_activity_verification_does_not_penalize_mastery(self):
        assignment = Assignment.objects.create(classroom=self.room, topic=self.topic, title="List", instructions="Insert", tests=[], require_camera=False)
        attempt = AssignmentAttempt.objects.create(assignment=assignment, student=self.student.student, code="head = new", status="submitted", test_score=100)
        MonitoringEvent.objects.create(attempt=attempt, event_id="observed", kind="page_hidden")
        response = self.post(self.sc, f"learning/courses/{self.course.id}/vivas/", {"assignment_attempt_id": attempt.id})
        self.assertTrue(response.json()["session"]["verification"])
        self.assertEqual(LearningEvidence.objects.count(), 0)
        self.assertNotIn("cheating", response.content.decode())
        session = response.json()["session"]
        response = self.post(self.sc, f"learning/vivas/{session['id']}/answer/", {"turn_id": session["turns"][0]["id"], "answer": "Connect the new node to the previous head before updating the head pointer"})
        self.assertEqual(len(response.json()["session"]["turns"]), 2)
        self.assertFalse(LearningEvidence.objects.get().verified)

    def test_insights_support_evidence_and_denominators(self):
        record_evidence(self.student.student, self.topic, "weak", "programming", 20, misconception="Pointer update")
        response = self.tc.get(f"/api/learning/courses/{self.course.id}/insights/")
        data = response.json()
        self.assertEqual(data["topics"][0]["assessed_students"], 1)
        self.assertIn("100%", data["observations"][0]["text"])
        self.assertTrue(data["observations"][0]["support"][0]["evidence_ids"])
        self.assertEqual(self.tc2.get(f"/api/learning/courses/{self.course.id}/insights/").status_code, 404)
        self.assertEqual(self.sc.get(f"/api/learning/courses/{self.course.id}/insights/").status_code, 403)
        self.assertEqual(self.tc.get(f"/api/learning/courses/{self.course.id}/insights/?format=csv").status_code, 200)

    def test_topic_detail_and_progress_are_student_scoped(self):
        self.upload()
        chunk = SourceChunk.objects.get()
        self.post(self.sc, f"learning/sources/{chunk.id}/", {"seconds_viewed": 30})
        report = self.sc.get(f"/api/learning/courses/{self.course.id}/report/").json()
        self.assertEqual(report["resources_studied"][0]["seconds_viewed"], 30)
        self.assertEqual(self.sc.get(f"/api/learning/courses/{self.course.id}/report/?student_id={self.other.student_id}").status_code, 403)
        self.assertEqual(self.tc.get(f"/api/learning/topics/{self.topic.id}/?student_id={self.student.student_id}").status_code, 200)

    def test_teacher_question_sources_cannot_cross_course(self):
        response = self.post(self.tc, f"learning/topics/{self.topic.id}/bank/", {"prompt": "Q", "options": ["A", "B"], "answer": {"correct_index": 0}, "source_ids": [999999]})
        self.assertEqual(response.status_code, 400)

    def test_demo_uses_real_sources_execution_and_saved_improvement(self):
        import ast, os, subprocess, sys, shutil, uuid
        from .learning.demo import INSERT_SOLUTION, COUNT_SOLUTION
        tree = ast.parse(Path(__file__).with_name("views.py").read_text())
        functions = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in ("run_tests", "run_python_code", "run_c_code", "run_java_code")]
        namespace = {"__name__": "labtwin.views", "subprocess": subprocess, "tempfile": tempfile, "os": os, "sys": sys, "shutil": shutil}
        exec(compile(ast.Module(body=functions, type_ignores=[]), "<trusted-demo-runners>", "exec"), namespace)
        with patch.object(self.fake_views, "run_tests", namespace["run_tests"], create=True):
            response = self.post(self.tc, "learning/demo/", {})
            self.assertEqual(response.status_code, 201, response.content)
            demo = response.json()["demo"]
            self.assertTrue(all(m["status"] == "ready" for m in demo["materials"]), demo)
            self.assertEqual(CourseMaterial.objects.count(), 3)
            self.assertEqual(self.post(self.tc, "learning/demo/", {}).json()["demo"]["id"], demo["id"])
            signed = self.post(self.tc, f"learning/demo/{demo['id']}/student/", {}).json()
            client = Client(HTTP_AUTHORIZATION="Bearer " + signed["token"])
            self.assertEqual(client.get(f"/api/learning/courses/{self.course.id}/materials/").status_code, 404)
            answer = self.post(client, f"learning/courses/{demo['course']['id']}/ask/", {"question": "Explain insertion at beginning of linked list"}).json()
            self.assertEqual({s["kind"] for s in answer["citations"]}, {"pdf", "slides", "video"})
            first = self.post(client, f"assignments/{demo['assignment_id']}/attempts/", {}).json()["attempt"]
            failed = self.post(client, f"attempts/{first['id']}/", {"code": first["code"]}).json()["attempt"]
            self.assertLess(failed["test_score"], 100)
            self.post(client, "learning/hint/", {"assignment_attempt_id": first["id"], "level": 2})
            retry = self.post(client, f"assignments/{demo['assignment_id']}/attempts/", {}).json()["attempt"]
            passed = self.post(client, f"attempts/{retry['id']}/", {"code": INSERT_SOLUTION}).json()["attempt"]
            self.assertEqual(passed["test_score"], 100)
            question = self.post(client, f"learning/courses/{demo['course']['id']}/practice/", {"topic_id": demo["topic_id"]}).json()["question"]
            result = self.post(client, f"learning/questions/{question['id']}/attempts/", {"request_id": str(uuid.uuid4()), "code": COUNT_SOLUTION})
            self.assertEqual(result.json()["attempt"]["score"], 100, result.content)
            insights = self.tc.get(f"/api/learning/courses/{demo['course']['id']}/insights/").json()
            self.assertTrue(any("improved" in row["text"] for row in insights["observations"]))
            self.assertEqual(len(insights["evidence"]), 4)
            self.assertEqual(sum(e.kind == "conversation" and e.score is None for e in LearningEvidence.objects.all()), 1)
            self.assertTrue(all(self.tc.get("/api/learning/demo/").json()["demo"]["progress"]))
            self.assertEqual(self.post(self.tc2, f"learning/demo/{demo['id']}/student/", {}).status_code, 404)

    def test_demo_disabled_by_default_configuration(self):
        with override_settings(LABTWIN_DEMO_ENABLED=False):
            self.assertEqual(self.post(self.tc, "learning/demo/", {}).status_code, 403)

    def test_teacher_review_supersedes_provisional_viva_evidence(self):
        session = self.post(self.sc, f"learning/courses/{self.course.id}/vivas/", {"topic_id": self.topic.id}).json()["session"]
        turn_id = session["turns"][0]["id"]
        self.post(self.sc, f"learning/vivas/{session['id']}/answer/", {"turn_id": turn_id, "answer": "Preserve the previous head connection"})
        self.assertEqual(self.post(self.sc, f"learning/viva-turns/{turn_id}/review/", {"score": 80}).status_code, 403)
        self.assertEqual(self.post(self.tc2, f"learning/viva-turns/{turn_id}/review/", {"score": 80}).status_code, 404)
        self.assertEqual(self.post(self.tc, f"learning/viva-turns/{turn_id}/review/", {"score": 90, "feedback": "Good reasoning"}).status_code, 200)
        self.assertEqual(self.post(self.tc, f"learning/viva-turns/{turn_id}/review/", {"score": 70}).status_code, 200)
        scored = LearningEvidence.objects.filter(score__isnull=False)
        self.assertEqual(scored.count(), 1)
        self.assertEqual(scored.get().score, 70)
        self.assertTrue(scored.get().verified)
        self.tc.delete(f"/api/classrooms/{self.room.id}/students/{self.student.student_id}/")
        self.assertEqual(self.post(self.tc, f"learning/viva-turns/{turn_id}/review/", {"score": 70}).status_code, 404)

    def test_inspect_actual_student_and_activity_evidence(self):
        from .learning.mastery import bridge_assignment
        assignment = Assignment.objects.create(classroom=self.room, topic=self.topic, title="Insert", instructions="Insert", tests=[], require_camera=False)
        attempt = AssignmentAttempt.objects.create(assignment=assignment, student=self.student.student, code="head = node", status="submitted", test_score=50)
        MonitoringEvent.objects.create(attempt=attempt, kind="page_hidden", event_id="obs")
        bridge_assignment(attempt)
        row = LearningEvidence.objects.get()
        endpoint = f"/api/learning/evidence/{row.id}/"
        self.assertEqual(self.sc.get(endpoint).status_code, 403)
        self.assertEqual(self.tc2.get(endpoint).status_code, 404)
        detail = self.tc.get(endpoint).json()["assignment_attempt"]
        self.assertEqual(detail["code"], "head = node")
        self.assertEqual(detail["event_counts"]["page_hidden"], 1)
        self.assertIn("not proof", detail["review_note"])

    def test_legacy_transport_hides_tests_without_erasing_server_history(self):
        from .models import StudentResponse
        row = StudentResponse.objects.create(student=self.student.student, stage="analysis", question="Q", test_results=[{"test": 1, "input": "hidden input", "expected": "hidden expected", "stdout": "hidden input", "passed": True}], result={"test_results": [{"input": "hidden input", "expected": "hidden expected", "passed": True}]})
        student_data = self.sc.get("/api/response-history/").json()["responses"][0]
        self.assertNotIn("hidden input", json.dumps(student_data))
        self.assertTrue(student_data["test_results"][0]["hidden"])
        teacher_data = self.tc.get(f"/api/response-history/?student_id={self.student.student_id}").json()["responses"][0]
        self.assertEqual(teacher_data["test_results"][0]["input"], "hidden input")
        row.refresh_from_db(); self.assertEqual(row.test_results[0]["input"], "hidden input")

    def test_automatic_transcription_preserves_segment_offsets(self):
        from .learning.transcription import transcribe
        from types import SimpleNamespace
        def convert(command, **kwargs):
            Path(command[-1]).write_bytes(b"test-audio")
        result = SimpleNamespace(model_dump=lambda: {"segments": [{"start": 1, "end": 8, "text": "Connect the new node to the old head."}]})
        with override_settings(LABTWIN_DISABLE_REMOTE_AI=False), patch.dict("os.environ", {"GROQ_API_KEY": "test-only-key"}), patch("labtwin.learning.transcription.media_info", return_value=1810), patch("labtwin.learning.transcription.subprocess.run", side_effect=convert), patch("groq.Groq") as provider:
            provider.return_value.audio.transcriptions.create.return_value = result
            rows, duration = transcribe("example.mp4")
        self.assertEqual([row["start"] for row in rows], [1, 901, 1801])
        self.assertEqual(duration, 1810)
        self.assertEqual(provider.return_value.audio.transcriptions.create.call_count, 3)
        self.assertEqual(provider.return_value.audio.transcriptions.create.call_args.kwargs["timestamp_granularities"], ["segment"])

    def test_failed_index_retains_file_and_retry_processes_again(self):
        with patch("labtwin.learning.ingestion.index_chunks", side_effect=RuntimeError("outage")):
            material = self.upload().json()["material"]
        self.assertEqual(material["status"], "failed")
        response = self.post(self.tc, f"learning/materials/{material['id']}/retry/", {})
        self.assertEqual(response.json()["material"]["status"], "ready", response.content)
        self.assertEqual(SourceChunk.objects.filter(unit__archived=False).count(), 1)
        self.assertEqual(SourceChunk.objects.filter(unit__archived=True).count(), 1)

    def test_grounded_ai_answer_validates_exact_quote_and_source(self):
        self.upload()
        chunk = SourceChunk.objects.get()
        with patch("labtwin.learning.rag.structured_completion", return_value={"supported": True, "claims": [{"text": "Preserve the old chain before changing head.", "quote": "connect the new node to the previous head", "source_id": chunk.id}]}):
            response = self.post(self.sc, f"learning/courses/{self.course.id}/ask/", {"question": "linked lists"}).json()
        self.assertEqual(response["mode"], "ai_grounded")
        self.assertEqual(response["citations"][0]["id"], chunk.id)

    def test_isolated_runner_contract_keeps_answers_in_application(self):
        from .execution import remote_tests
        import contextlib
        result = io.BytesIO(json.dumps({"executions": [{"success": True, "stdout": "5", "stderr": ""}]}).encode())
        with override_settings(LABTWIN_RUNNER_URL="https://private.example", LABTWIN_RUNNER_SECRET="test-only-secret"), patch("labtwin.execution.urlopen", return_value=contextlib.closing(result)) as service:
            score, rows = remote_tests("print(5)", [{"input": "2 3", "expected": "5"}], "Python")
        self.assertEqual(score, 100)
        self.assertNotIn("expected", service.call_args.args[0].data.decode())
        self.assertEqual(service.call_args.args[0].get_header("Authorization"), "Bearer test-only-secret")
        self.assertEqual(rows[0]["expected"], "5")
