"""Focused failure-injection checks; these do not certify live AI providers."""
import io
import os
import subprocess
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import RequestFactory, SimpleTestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile

from .models import Course, CourseMaterial, SourceChunk
from .test_classrooms import ClassroomFixture
from .learning.access import LearningError
from .learning.extraction import extract, pdf_units, slide_units


class ExtractionReliabilityTests(SimpleTestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(prefix="labtwin-reliability-")
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)

    def image(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (1200, 300), "white")
        ImageDraw.Draw(image).text((20, 45), "Connect the new node before updating head.",
            font=ImageFont.truetype("DejaVuSans.ttf", 44), fill="black")
        return image

    def pdf(self, text="", image=True):
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas
        path = self.root / "notes.pdf"
        writer = canvas.Canvas(str(path), pagesize=(600, 500))
        if text:
            writer.drawString(25, 450, text)
        if image:
            writer.drawImage(ImageReader(self.image()), 0, 100, 600, 150)
        writer.showPage(); writer.save()
        return path

    def slides(self, text):
        from pptx import Presentation
        from pptx.util import Inches
        image_path = self.root / "figure.png"; self.image().save(image_path)
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(.2), Inches(.2), Inches(9), Inches(1))
        box.text = text
        slide.shapes.add_picture(str(image_path), Inches(.2), Inches(1.3), width=Inches(9))
        path = self.root / "slides.pptx"; deck.save(path)
        return path

    def test_native_text_pdf_does_not_require_ocr(self):
        path = self.pdf("Linked lists connect nodes through pointers.", image=False)
        with patch("labtwin.learning.extraction.shutil.which", return_value=None):
            units, _ = pdf_units(path)
        self.assertIn("connect nodes", units[0]["text"])
        self.assertEqual(units[0]["page_number"], 1)

    def test_sparse_headers_do_not_make_scanned_pdf_ready_without_ocr(self):
        path = self.pdf("Program: Output:")
        with patch("labtwin.learning.extraction.shutil.which", return_value=None):
            with self.assertRaisesRegex(LearningError, "scanned PDF"):
                pdf_units(path)

    def test_scanned_pdf_uses_real_installed_ocr(self):
        units, _ = pdf_units(self.pdf())
        self.assertTrue(any("updating head" in row["text"] for row in units))
        self.assertTrue(all(row["page_number"] == 1 for row in units))

    def test_optional_figure_ocr_timeout_preserves_native_teaching_text(self):
        path = self.pdf("Linked lists store data in nodes connected by pointers. Preserve the old head before updating it.")
        with patch("labtwin.learning.extraction.image_text", side_effect=subprocess.TimeoutExpired("tesseract", 30)):
            units, warnings = pdf_units(path)
        self.assertTrue(any("Preserve the old head" in row["text"] for row in units))
        self.assertTrue(any(row.get("visual_blob") for row in units))
        self.assertTrue(any("OCR" in warning for warning in warnings))

    def test_required_ocr_timeout_is_recoverable_not_corrupted_pdf(self):
        with patch("labtwin.learning.extraction.image_text", side_effect=subprocess.TimeoutExpired("tesseract", 30)):
            with self.assertRaisesRegex(LearningError, "scanned PDF"):
                pdf_units(self.pdf())

    def test_slide_ocr_timeout_keeps_image_text_and_source_location(self):
        path = self.slides("Linked lists store data in nodes and connect them with pointers.")
        with patch("labtwin.learning.extraction.image_text", side_effect=subprocess.TimeoutExpired("tesseract", 30)):
            units, warnings = slide_units(path)
        self.assertIn("connect them", units[0]["text"])
        figures = [row for row in units if row.get("content_type") == "visual"]
        self.assertTrue(figures, "A failed optional OCR must not discard image pixels.")
        self.assertEqual(figures[0]["slide_number"], 1)
        self.assertTrue(figures[0]["visual_blob"])
        self.assertTrue(any("OCR" in warning for warning in warnings))

    def test_unanalysed_slide_image_is_not_indexed_as_a_fake_description(self):
        with patch("labtwin.learning.extraction.shutil.which", return_value=None):
            units, warnings = slide_units(self.slides("Linked lists connect nodes."))
        figure = next(row for row in units if row.get("content_type") == "visual")
        self.assertEqual(figure["text"], "")
        self.assertEqual(figure["visual_description"], "")
        self.assertEqual(figure["layout"]["analysis"]["semantic_status"], "limited")
        self.assertTrue(warnings)

    def test_native_slide_table_remains_searchable_without_vision(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(3)).table
        for cell, text in zip([table.cell(0,0), table.cell(0,1), table.cell(1,0), table.cell(1,1)], ["Structure", "Nodes", "Linked list", "3"]):
            cell.text = text
        path = self.root / "table.pptx"; deck.save(path)
        units, _ = slide_units(path)
        self.assertIn("Linked list | 3", units[0]["text"])

    @override_settings(LABTWIN_VIDEO_OCR=False)
    def test_real_embedded_captions_are_used_before_speech_provider(self):
        path = Path(__file__).parent / "demo_assets" / "Lecture-4.webm"
        material = SimpleNamespace(kind="video", file=SimpleNamespace(path=str(path)), course=SimpleNamespace(is_demo=False))
        with patch("labtwin.learning.extraction.transcribe") as speech:
            units, _, extra, _ = extract(material)
        speech.assert_not_called()
        self.assertEqual(extra["transcript_origin"], "embedded_captions")
        self.assertTrue(any(row["start_seconds"] > 0 for row in units))

    def test_speech_unavailable_has_defined_safe_error(self):
        from .learning.transcription import transcribe
        with patch.dict(os.environ, {"GROQ_API_KEY": ""}):
            with self.assertRaisesRegex(LearningError, "transcription"):
                transcribe(Path(__file__).parent / "demo_assets" / "Lecture-4.webm")

    @override_settings(LABTWIN_DISABLE_REMOTE_AI=False)
    def test_speech_provider_timeout_is_bounded_and_redacted(self):
        from .learning.transcription import transcribe
        provider = MagicMock()
        provider.audio.transcriptions.create.side_effect = TimeoutError("private-provider-response")
        with patch.dict(os.environ, {"GROQ_API_KEY": "not-a-real-key"}), patch("groq.Groq", return_value=provider) as constructor:
            with self.assertRaises(LearningError) as raised:
                transcribe(Path(__file__).parent / "demo_assets" / "Lecture-4.webm")
        self.assertNotIn("private-provider-response", raised.exception.message)
        self.assertIn("upload is retained", raised.exception.message)
        self.assertEqual(constructor.call_args.kwargs["timeout"], 120)
        self.assertEqual(constructor.call_args.kwargs["max_retries"], 1)


class ProviderReliabilityTests(SimpleTestCase):
    def test_original_coach_error_response_and_logs_do_not_reveal_private_provider_data(self):
        from . import views
        request = RequestFactory().post("/api/tutor/", {"concept_key": "Pointers"}, content_type="application/json")
        with patch("labtwin.views.get_current_question", side_effect=RuntimeError("private-provider-response /internal/private/path")), patch("builtins.print") as printer:
            response = views.tutor_help(request)
        self.assertEqual(response.status_code, 500)
        self.assertNotIn("private-provider-response", response.content.decode())
        self.assertNotIn("/internal/private/path", response.content.decode())
        self.assertNotIn("private-provider-response", str(printer.call_args_list))

    @override_settings(LABTWIN_DISABLE_REMOTE_AI=False)
    def test_tutoring_provider_failure_returns_grounded_fallback_without_leaking_error(self):
        from .learning.ai import structured_completion
        provider = MagicMock()
        provider.chat.completions.create.side_effect = TimeoutError("private-provider-response")
        with patch.dict(os.environ, {"GROQ_API_KEY": "not-a-real-key"}), patch("groq.Groq", return_value=provider) as constructor:
            with self.assertLogs("labtwin.learning.ai", level="WARNING") as logs:
                result = structured_completion("Use the source", {"source": "Head points to the first node."})
        self.assertIsNone(result)
        self.assertNotIn("private-provider-response", " ".join(logs.output))
        self.assertEqual(constructor.call_args.kwargs["timeout"], 35)
        self.assertEqual(constructor.call_args.kwargs["max_retries"], 1)

    @override_settings(LABTWIN_DISABLE_REMOTE_AI=False, LABTWIN_VISION_MODEL="configured-vision-model")
    def test_vision_timeout_keeps_native_provenance_and_actual_pixels(self):
        from .learning.visuals import interpret_image
        provider = MagicMock()
        provider.chat.completions.create.side_effect = TimeoutError("private-provider-response")
        graph = {"nodes": [{"id": "a", "label": "Head"}, {"id": "b", "label": "Node"}],
                 "edges": [{"from": "a", "to": "b", "directed": True}]}
        with patch.dict(os.environ, {"GROQ_API_KEY": "not-a-real-key"}), patch("groq.Groq", return_value=provider) as constructor:
            analysis = interpret_image(b"actual-image-pixels", "Nearby source caption", graph)
        self.assertEqual(analysis["method"], "native_geometry")
        self.assertEqual(analysis["description"], "Head points to Node.")
        self.assertIn("unavailable", analysis["warning"])
        self.assertNotIn("private-provider-response", str(analysis))
        request = provider.chat.completions.create.call_args.kwargs
        self.assertIn("YWN0dWFsLWltYWdlLXBpeGVscw==", request["messages"][0]["content"][1]["image_url"]["url"])
        self.assertEqual(constructor.call_args.kwargs["timeout"], 45)
        self.assertEqual(constructor.call_args.kwargs["max_retries"], 1)

    def test_original_programming_coach_has_a_provider_response_timeout(self):
        from .views import llm
        self.assertIsNotNone(llm.timeout, "The original programming AI path also needs a provider timeout.")
        self.assertGreater(llm.timeout, 0)
        self.assertLessEqual(llm.timeout, 120)


@override_settings(LABTWIN_PROCESS_INLINE=True, LABTWIN_PROCESS_MODE="external", LABTWIN_VIDEO_OCR=False)
class UploadReliabilityTests(ClassroomFixture):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.folder = tempfile.TemporaryDirectory(prefix="labtwin-upload-reliability-")
        cls.storage = override_settings(LABTWIN_MEDIA_ROOT=cls.folder.name + "/uploads", LABTWIN_VECTOR_ROOT=cls.folder.name + "/vectors")
        cls.storage.enable()

    @classmethod
    def tearDownClass(cls):
        cls.storage.disable()
        super().tearDownClass()
        cls.folder.cleanup()

    def setUp(self):
        super().setUp(); self.join()
        self.course = Course.objects.create(classroom=self.room, name="Data Structures")

    def upload(self, name, blob):
        response = self.tc.post(f"/api/learning/courses/{self.course.id}/materials/", {"file": SimpleUploadedFile(name, blob)})
        self.assertEqual(response.status_code, 201, response.content)
        return CourseMaterial.objects.get(pk=response.json()["material"]["id"])

    def raster_pdf(self, native_text):
        from pypdf import PdfReader, PdfWriter
        from reportlab.pdfgen import canvas
        from .test_material_processing import MaterialProcessingTests
        header = io.BytesIO(); drawing = canvas.Canvas(header, pagesize=(600, 225))
        drawing.drawString(20, 210, native_text); drawing.showPage(); drawing.save()
        page = PdfReader(io.BytesIO(MaterialProcessingTests.scan_pdf())).pages[0]
        page.merge_page(PdfReader(header).pages[0])
        writer = PdfWriter(); writer.add_page(page)
        blob = io.BytesIO(); writer.write(blob)
        return blob.getvalue()

    def test_missing_required_ocr_upload_terminates_failed_with_retained_original(self):
        with patch("labtwin.learning.extraction.shutil.which", return_value=None):
            material = self.upload("header-only-scan.pdf", self.raster_pdf("Program: Output:"))
        self.assertEqual(material.status, "failed")
        self.assertEqual(material.job.status, "failed")
        self.assertIn("scanned PDF", material.error)
        self.assertTrue(material.file.storage.exists(material.file.name))
        self.assertFalse(SourceChunk.objects.filter(course=self.course).exists())

    def test_optional_ocr_failure_upload_still_finishes_ready_with_warning(self):
        blob = self.raster_pdf("Linked lists preserve the previous head. Connect the new node before updating the head pointer.")
        with patch("labtwin.learning.extraction.image_text", side_effect=subprocess.TimeoutExpired("tesseract", 30)):
            material = self.upload("native-plus-figure.pdf", blob)
        self.assertEqual(material.status, "ready", material.error)
        self.assertEqual(material.job.status, "ready")
        self.assertTrue(material.warnings)
        self.assertTrue(material.units.filter(content_type="visual").exists())
        self.assertTrue(SourceChunk.objects.filter(course=self.course, text__icontains="previous head").exists())

    def test_real_captionless_video_without_speech_service_terminates_failed(self):
        original = Path(__file__).parent / "demo_assets" / "Lecture-4.webm"
        video = Path(self.folder.name) / "captionless.webm"
        subprocess.run(["ffmpeg", "-y", "-nostdin", "-v", "error", "-i", str(original), "-map", "0:v:0", "-map", "0:a:0", "-t", "2", "-c", "copy", "-sn", str(video)], capture_output=True, check=True, timeout=30)
        material = self.upload("captionless.webm", video.read_bytes())
        self.assertEqual(material.status, "failed")
        self.assertEqual(material.job.status, "failed")
        self.assertIn("transcription", material.error)
        self.assertTrue(material.file.storage.exists(material.file.name))
