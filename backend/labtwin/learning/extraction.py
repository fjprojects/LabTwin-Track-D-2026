"""Format-specific extraction keeps source boundaries intact."""
import base64
import io
import shutil
import subprocess
import tempfile
import zipfile
import math
import logging
import re
from pathlib import Path
from django.conf import settings

from .access import LearningError
from .transcription import media_info, transcribe, validate_transcript

KINDS = {".pdf": "pdf", ".pptx": "slides", ".ppt": "slides", ".txt": "text", ".md": "text", ".mp4": "video", ".webm": "video", ".mov": "video", ".mkv": "video", ".mp3": "audio", ".wav": "audio", ".m4a": "audio", ".ogg": "audio", ".flac": "audio"}
logger = logging.getLogger(__name__)


def image_text(blob):
    if not shutil.which("tesseract"):
        return ""
    with tempfile.TemporaryDirectory(prefix="labtwin-image-ocr-") as temp:
        path = Path(temp) / "image.jpg"
        path.write_bytes(blob)
        result = subprocess.run(["tesseract", str(path), "stdout"], capture_output=True, timeout=30, check=True)
        return result.stdout.decode("utf-8", errors="replace").strip()


def safe_image_text(blob, warnings, location):
    """An optional figure OCR failure must not discard native teaching text."""
    try:
        return image_text(blob)
    except (subprocess.SubprocessError, OSError) as error:
        warnings.append(f"{location}: OCR was unavailable; the image and any native text are retained.")
        # Do not log parser/provider messages, source contents, or local paths.
        logger.warning("Optional image OCR unavailable (%s)", type(error).__name__)
        return ""


def video_frames(path, duration):
    if not getattr(settings, "LABTWIN_VIDEO_OCR", True) or not shutil.which("tesseract"):
        return [], []
    interval = max(30, math.ceil(duration / 80))
    rows = []
    with tempfile.TemporaryDirectory(prefix="labtwin-video-ocr-") as temp:
        for offset in range(0, math.ceil(duration), interval):
            image = Path(temp) / "frame.jpg"
            try:
                subprocess.run(["ffmpeg", "-y", "-nostdin", "-v", "error", "-ss", str(offset), "-i", str(path), "-frames:v", "1", "-vf", "scale=1280:-1", str(image)], capture_output=True, timeout=30, check=True)
                text = image_text(image.read_bytes())
                if len(text) >= 30:
                    rows.append({"text": text, "start_seconds": offset, "end_seconds": min(duration, offset + interval), "layout": {"origin": "video_frame_ocr"}})
            except (subprocess.SubprocessError, OSError):
                return rows, ["Some video frames could not be OCR-processed. Speech/caption sources remain available."]
    return rows, [f"Visual text was sampled every {interval} seconds. Frame OCR does not represent every visual event."] if rows else []


def convert_office(path, extension, directory):
    if not shutil.which("libreoffice"):
        raise LearningError("Legacy PPT conversion requires LibreOffice. Save the slides as PPTX and upload again.")
    profile = Path(directory) / "office-profile"
    try:
        subprocess.run(["libreoffice", f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", extension, "--outdir", directory, str(path)], capture_output=True, timeout=120, check=True)
        output = Path(directory) / (Path(path).stem + "." + extension)
        if not output.exists():
            raise LearningError("The presentation could not be converted.")
        return output
    except subprocess.SubprocessError as exc:
        raise LearningError("This presentation could not be read. Try exporting it to PPTX or PDF.") from exc


def pdf_units(path, progress=None):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise LearningError("PDF extraction is not installed. Install requirements-render.txt in the Python environment running LabTwin.") from exc
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise LearningError("Password-protected PDFs are unsupported. Upload an unlocked copy.")
        if len(reader.pages) > 1000:
            raise LearningError("Upload a PDF with at most 1,000 pages.")
        units, warnings = [], []
        for number, page in enumerate(reader.pages, 1):
            if progress:
                progress("pdf_text", number - 1, len(reader.pages), "pages")
            text = (page.extract_text() or "").strip()
            # Native headers such as 'Program: Output:' are not the scanned
            # teaching content. Keep them, but do not mistake them for OCR.
            needs_ocr = not text or (len(re.findall(r"\w+", text)) < 8 and bool(len(page.images)))
            ocr_complete = False
            if needs_ocr and shutil.which("tesseract"):
                if progress:
                    progress("pdf_ocr", number - 1, len(reader.pages), "pages")
                # PyMuPDF is already a project dependency. Rendering here also
                # works on Windows without a separately installed Poppler.
                import fitz
                with fitz.open(path) as document:
                    if document.is_encrypted:
                        document.authenticate("")
                    scanned_page = document[number - 1]
                    scale = min(3, 1800 / max(scanned_page.rect.width, scanned_page.rect.height))
                    recognised = safe_image_text(scanned_page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes("png"), warnings, f"Page {number}")
                    if recognised:
                        text, ocr_complete = recognised, True
            elif needs_ocr:
                logger.warning("Scanned PDF page needs OCR; Tesseract is unavailable on this deployment")
            if text:
                units.append({"number": number, "page_number": number, "text": text,
                              "layout": {"origin": "pdf_ocr" if ocr_complete else "pdf_native_text", "requires_ocr": needs_ocr and not ocr_complete}})
                if needs_ocr and not ocr_complete:
                    warnings.append(f"Page {number}: native headers were preserved, but scanned teaching content could not be read.")
            else:
                warnings.append(f"Page {number} has no extractable text.")
            if progress:
                progress("pdf_text", number, len(reader.pages), "pages")
        from .visuals import pdf_visual_units
        visual_units, visual_warnings = pdf_visual_units(path, {u["page_number"]: u["text"] for u in units}, progress=progress)
        units += visual_units
        if not any(unit["text"].strip() and not unit.get("layout", {}).get("requires_ocr") for unit in units):
            raise LearningError("This appears to be a scanned PDF. No readable teaching content was found: OCR is required, but usable OCR or visual extraction is not available on this deployment. Ask the administrator to enable Tesseract OCR, or upload a searchable PDF. The original file is retained.")
        units.sort(key=lambda item: (item["page_number"], item.get("content_type") == "visual"))
        for index, unit in enumerate(units, 1):
            unit["number"] = index
            if unit.get("content_type") != "visual":
                unit["title"] = unit["text"].split("\n")[0][:200]
        return units, warnings + visual_warnings
    except LearningError:
        raise
    except Exception as exc:
        raise LearningError("This PDF is corrupted or could not be extracted.") from exc


def slide_units(path, progress=None):
    from pptx import Presentation
    from PIL import Image
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 20000 or sum(e.file_size for e in entries) > 300 * 1024 * 1024:
                raise LearningError("The expanded presentation exceeds the processing limit.")
        deck = Presentation(path)
        if len(deck.slides) > 500:
            raise LearningError("Upload at most 500 slides.")
        units, warnings = [], []
        for number, slide in enumerate(deck.slides, 1):
            if progress:
                progress("slides", number - 1, len(deck.slides), "slides")
            texts, shapes, pictures = [], [], []
            for shape in slide.shapes:
                item = {"x": shape.left / deck.slide_width * 100, "y": shape.top / deck.slide_height * 100, "w": shape.width / deck.slide_width * 100, "h": shape.height / deck.slide_height * 100}
                if shape.has_text_frame:
                    texts.append(shape.text)
                    item["text"] = shape.text
                elif shape.has_table:
                    table_text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                    texts.append(table_text)
                    item["text"] = table_text
                elif shape.shape_type == 13:
                    try:
                        with Image.open(io.BytesIO(shape.image.blob)) as picture:
                            picture.thumbnail((1200, 1200))
                            buffer = io.BytesIO()
                            picture.convert("RGB").save(buffer, format="JPEG", quality=80)
                            item["image"] = "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()
                            visual_text = safe_image_text(buffer.getvalue(), warnings, f"Slide {number}")
                            if visual_text:
                                texts.append("Slide image text: " + visual_text)
                            visual_buffer = io.BytesIO()
                            picture.convert("RGB").save(visual_buffer, format="PNG")
                            pictures.append((visual_buffer.getvalue(), visual_text))
                    except Exception:
                        warnings.append(f"Slide {number}: an image could not be decoded; native slide text remains available.")
                if "text" in item or "image" in item:
                    shapes.append(item)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                texts.append(slide.notes_slide.notes_text_frame.text)
            units.append({"number": number, "slide_number": number, "title": (slide.shapes.title.text if slide.shapes.title else f"Slide {number}")[:200], "text": "\n".join(texts).strip(), "layout": {"ratio": deck.slide_width / deck.slide_height, "shapes": shapes}})
            from .visuals import slide_graph, render_slide_graph, interpret_image
            graph = slide_graph(slide)
            if graph["edges"]:
                pictures.append((render_slide_graph(graph, deck.slide_width, deck.slide_height), ""))
            for blob, labels in pictures[:8]:
                analysis = interpret_image(blob, "\n".join(texts), graph if graph["edges"] else {})
                if not analysis["description"]:
                    if labels:
                        analysis["description"] = "Figure labels: " + labels
                    else:
                        warnings.append(f"Slide {number}: image preserved, but visual interpretation was unavailable. Readable slide text remains available.")
                if analysis.get("warning"):
                    warnings.append(analysis["warning"])
                units.append({"slide_number": number, "title": (slide.shapes.title.text if slide.shapes.title else f"Slide {number}")[:200],
                              "content_type": "visual", "text": "Educational diagram/figure. " + analysis["description"] if analysis["description"] else "",
                              "visual_description": analysis["description"], "analysis_method": analysis["method"],
                              "visual_blob": blob, "layout": {"origin": "slide_figure", "analysis": analysis, "nearby_text": "\n".join(texts)[:4000]}})
            if progress:
                progress("slides", number, len(deck.slides), "slides")
        if not any(u["text"] for u in units):
            raise LearningError("No searchable slide text was found. Add speaker notes or upload a PDF with OCR.")
        for index, unit in enumerate(units, 1):
            unit["number"] = index
        return units, warnings
    except LearningError:
        raise
    except Exception as exc:
        raise LearningError("The PPTX is corrupted or unsupported.") from exc


def extract(material, captions=None, progress=None):
    path = Path(material.file.path)
    if material.kind == "pdf":
        units, warnings = pdf_units(path, progress=progress)
        return units, warnings, {}, None
    if material.kind == "slides":
        with tempfile.TemporaryDirectory(prefix="labtwin-slides-") as temp:
            converted = convert_office(path, "pptx", temp) if path.suffix == ".ppt" else path
            units, warnings = slide_units(converted, progress=progress)
        return units, warnings, {}, None
    if material.kind == "text":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeError as exc:
            raise LearningError("Text materials must use UTF-8 encoding.") from exc
        if "\x00" in text or not text.strip():
            raise LearningError("This file is empty or is not plain text.")
        sections = [text[i:i + 12000] for i in range(0, len(text), 12000)]
        return [{"number": i + 1, "text": section} for i, section in enumerate(sections)], [], {}, None
    if progress:
        progress("transcribing")
    if captions is not None:
        duration = media_info(path)
        transcript = validate_transcript(captions, duration)
        origin = "teacher_captions"
    else:
        from .transcription import embedded_transcript
        embedded = embedded_transcript(path) if material.kind == "video" else None
        if embedded:
            transcript, duration = embedded
            origin = "embedded_captions"
        else:
            transcript, duration = transcribe(path)
            origin = "automatic_speech"
    units = [{"text": row["text"], "start_seconds": row["start"], "end_seconds": row["end"], "layout": {"origin": origin}} for row in transcript]
    if progress:
        progress("video_visuals")
    frames, warnings = video_frames(path, duration) if material.kind == "video" and not material.course.is_demo else ([], [])
    units = sorted(units + frames, key=lambda row: row["start_seconds"])
    for i, unit in enumerate(units, 1):
        unit["number"] = i
    return units, warnings, {"transcript": transcript, "transcript_origin": origin, "duration_seconds": duration}, None
